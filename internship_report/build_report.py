# -*- coding: utf-8 -*-
"""Build the FieldSync internship report as a Word (.docx) document."""

import os

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import report_content_1 as c1
import report_content_2 as c2
import appendix_content as ap

BASE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(BASE, "assets")
OUT = os.path.join(BASE, "DBU_Internship_Report_FieldSync_2018_EC.docx")

BLACK = RGBColor(0x00, 0x00, 0x00)
TNR = "Times New Roman"


# ---------------------------------------------------------------- helpers
def _rfonts(rpr, name):
    r = rpr.get_or_add_rFonts()
    r.set(qn("w:ascii"), name)
    r.set(qn("w:hAnsi"), name)
    r.set(qn("w:cs"), name)
    r.set(qn("w:eastAsia"), name)


def set_run(run, name=TNR, size=12, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    _rfonts(run._element.get_or_add_rPr(), name)


def set_pg_fmt(section, fmt, start=None):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    if pg.getparent() is None:
        pgMar = sectPr.find(qn("w:pgMar"))
        if pgMar is not None:
            pgMar.addnext(pg)
        else:
            sectPr.insert(0, pg)
    return pg


def add_field(paragraph, code, placeholder=""):
    run = paragraph.add_run()
    r = run._r
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); b.set(qn("w:dirty"), "true")
    r.append(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code
    r.append(it)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r.append(s)
    r2 = paragraph.add_run(placeholder)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r2._r.append(e)
    return r2


def add_page_field(paragraph):
    add_field(paragraph, "PAGE")


def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def shade_paragraph(paragraph, hexcolor):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    pPr.append(shd)


def border_paragraph(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + side)
        el.set(qn("w:val"), "dashed")
        el.set(qn("w:sz"), "12")
        el.set(qn("w:space"), "8")
        el.set(qn("w:color"), "6B7280")
        pBdr.append(el)
    pPr.append(pBdr)


def add_seq_caption(paragraph, label):
    run = paragraph.add_run()
    r = run._r
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); b.set(qn("w:dirty"), "true")
    r.append(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = "SEQ %s \\* ARABIC" % label
    r.append(it)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r.append(s)
    r2 = paragraph.add_run("0")
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r2._r.append(e)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = TNR
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    _rfonts(normal.element.get_or_add_rPr(), TNR)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def _heading(name, size, bold=True, center=False, space_before=12, space_after=6, italic=False):
        st = doc.styles[name]
        st.font.name = TNR
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = italic
        st.font.color.rgb = BLACK
        _rfonts(st.element.get_or_add_rPr(), TNR)
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(space_after)
        st.paragraph_format.line_spacing = 1.2
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        # keep_with_next so headings are not orphaned
        st.paragraph_format.keep_with_next = True
        return st

    _heading("Heading 1", 14, center=True, space_before=6, space_after=14)
    _heading("Heading 2", 13, space_before=14, space_after=6)
    _heading("Heading 3", 12, space_before=12, space_after=6, italic=True)

    # CodeBlock
    cb = doc.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
    cb.base_style = doc.styles["Normal"]
    cb.font.name = "Consolas"
    cb.font.size = Pt(9)
    cb.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    rpr = cb.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:ascii"), "Consolas")
    rpr.get_or_add_rFonts().set(qn("w:hAnsi"), "Consolas")
    cb.paragraph_format.line_spacing = 1.0
    cb.paragraph_format.space_after = Pt(0)
    cb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Caption
    try:
        cap = doc.styles["Caption"]
    except KeyError:
        cap = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
        cap.base_style = doc.styles["Normal"]
    cap.font.name = TNR
    cap.font.size = Pt(10.5)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    _rfonts(cap.element.get_or_add_rPr(), TNR)
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.line_spacing = 1.0
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def setup_sections(doc):
    secs = doc.sections
    for s in secs:
        s.page_width = Mm(210)
        s.page_height = Mm(297)
        for a in ("left_margin", "right_margin"):
            setattr(s, a, Inches(1))
        for a in ("top_margin", "bottom_margin"):
            setattr(s, a, Inches(1))
    return doc.sections[0]


# ---------------------------------------------------------------- renderers
def para(doc, text, style=None, justify=True, hanging=False):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run(r)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if hanging:
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(8)
    return p


def render_blocks(doc, blocks):
    for kind, payload in blocks:
        if kind == "h1":
            p = doc.add_paragraph(style="Heading 1")
            r = p.add_run(payload)
            set_run(r, size=14, bold=True)
            p.paragraph_format.page_break_before = True
        elif kind == "h2":
            p = doc.add_paragraph(style="Heading 2")
            r = p.add_run(payload)
            set_run(r, size=13, bold=True)
        elif kind == "h3":
            p = doc.add_paragraph(style="Heading 3")
            r = p.add_run(payload)
            set_run(r, size=12, bold=True, italic=True)
        elif kind == "p":
            para(doc, payload)
        elif kind == "bullets":
            for item in payload:
                p = doc.add_paragraph(style="List Bullet")
                r = p.add_run(item)
                set_run(r)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "numbered":
            for item in payload:
                p = doc.add_paragraph(style="List Number")
                r = p.add_run(item)
                set_run(r)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "ref":
            para(doc, payload, hanging=True)
        elif kind == "fig":
            fname = payload["file"]
            caption = payload["caption"]
            width = payload.get("width_in", 5.8)
            label = payload.get("label", "Figure")
            try:
                doc.add_picture(os.path.join(ASSETS, fname), width=Inches(width))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as exc:
                para(doc, "[missing image: %s]" % fname)
                para(doc, "")
            pcap = doc.add_paragraph(style="Caption")
            r1 = pcap.add_run(label + " ")
            set_run(r1, size=10.5, italic=True, color=RGBColor(0x55, 0x55, 0x55))
            add_seq_caption(pcap, label)
            r2 = pcap.add_run(": " + caption)
            set_run(r2, size=10.5, italic=True, color=RGBColor(0x55, 0x55, 0x55))
            pcap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "table":
            meta = payload
            cap = meta.get("caption", "")
            label = meta.get("label", "Table")
            headers = meta["headers"]
            rows = meta["rows"]
            widths = meta.get("widths")
            pcap = doc.add_paragraph(style="Caption")
            r1 = pcap.add_run(label + " ")
            set_run(r1, size=10.5, italic=True, color=RGBColor(0x55, 0x55, 0x55))
            add_seq_caption(pcap, label)
            r2 = pcap.add_run(": " + cap)
            set_run(r2, size=10.5, italic=True, color=RGBColor(0x55, 0x55, 0x55))
            pcap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            t = doc.add_table(rows=1 + len(rows), cols=len(headers))
            t.style = "Table Grid"
            for j, h in enumerate(headers):
                cell = t.cell(0, j)
                cell.text = ""
                p = cell.paragraphs[0]
                r = p.add_run(h)
                set_run(r, size=10.5, bold=True)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shade_cell(cell, "E5E7EB")
            for i, row in enumerate(rows, start=1):
                for j, val in enumerate(row):
                    cell = t.cell(i, j)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    r = p.add_run(str(val))
                    set_run(r, size=10.5)
            if widths:
                for j, w in enumerate(widths):
                    for i in range(len(t.rows)):
                        t.cell(i, j).width = Inches(w)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif kind == "code":
            for line in payload:
                p = doc.add_paragraph(style="CodeBlock")
                r = p.add_run(line if line else " ")
                set_run(r, name="Consolas", size=9, color=RGBColor(0x1F, 0x29, 0x37))
                shade_paragraph(p, "F3F4F6")
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
        elif kind == "appendix_h1":
            p = doc.add_paragraph(style="Heading 1")
            r = p.add_run(payload)
            set_run(r, size=14, bold=True)
            p.paragraph_format.page_break_before = True
        elif kind == "placeholder_box":
            p = doc.add_paragraph()
            border_paragraph(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(payload.replace("\n", " "))
            set_run(r, size=11, italic=True, color=RGBColor(0x6B, 0x72, 0x80))
            p.paragraph_format.space_after = Pt(24)


# ---------------------------------------------------------------- cover
def build_cover(doc):
    def line(text, size=12, bold=False, italic=False, space_after=4, spacer=0):
        for _ in range(spacer):
            doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, italic=italic)
        p.paragraph_format.space_after = Pt(space_after)
        return p

    line("DEBRE BERHAN UNIVERSITY", 22, bold=True, space_after=2)
    line("Debre Berhan, Ethiopia", 12, italic=True, space_after=18)
    line(c1.FACULTY, 16, bold=True, space_after=2)
    line(c1.DEPT, 16, bold=True, space_after=6)
    line("", 1, spacer=2)
    line("INTERNSHIP REPORT", 20, bold=True, space_after=4)
    line("On the Work Placement Undertaken at", 13, space_after=4)
    line(c1.COMPANY_FULL, 14, bold=True, space_after=6)
    line("Project: Development of the FieldSync Offline-First Registration",
         13, bold=True, space_after=2)
    line("and Reporting System", 13, bold=True, space_after=16)
    line("A report submitted in partial fulfilment of the requirements for the Degree",
         12, italic=True, space_after=2)
    line("of Bachelor of Science in Information Technology", 12, italic=True, space_after=16)

    details = [
        ("Student Name", c1.STUDENT),
        ("Student ID", c1.STUDENT_ID),
        ("Hosting Company", c1.COMPANY),
        ("Department", c1.DEPT),
        ("Duration of Internship", c1.DURATION),
        ("Date of Report Submission", c1.SUBMIT),
    ]
    t = doc.add_table(rows=len(details), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(details):
        c0 = t.cell(i, 0)
        c1c = t.cell(i, 1)
        for j, (c, txt, bold) in enumerate([(c0, k + ":", True), (c1c, v, False)]):
            cell = c
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(txt)
            set_run(r, size=12, bold=bold)
    t.columns[0].width = Inches(2.6)
    t.columns[1].width = Inches(4.0)
    for row in t.rows:
        row.cells[0].width = Inches(2.6)
        row.cells[1].width = Inches(4.0)

    line("", 1, spacer=2)
    line("Addis Ababa, Ethiopia", 12, space_after=2)


# ---------------------------------------------------------------- prelim pages
def build_declaration(doc):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("DECLARATION")
    set_run(r, size=14, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    body = (
        "I, %s, hereby declare that this internship report entitled "
        "\u201cDevelopment of the FieldSync Offline-First Registration and Reporting System\u201d is my "
        "original work and that no part of it has been previously presented for a degree or similar award "
        "in this or any other university. All sources of information used in this report are acknowledged "
        "through references.") % c1.STUDENT
    para(doc, body)
    para(doc, "The report has been prepared based on the internship work I carried out at AFRICOM "
              "Technologies PLC during the 2018 E.C. academic year, as partial fulfilment of the "
              "requirements for the Degree of Bachelor of Science in Information Technology at Debre "
              "Berhan University.")
    para(doc, "I declare that the information presented in this report is true and correct to the best of "
              "my knowledge and belief.")

    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("APPROVAL")
    set_run(r, size=13, bold=True)
    p.paragraph_format.space_after = Pt(8)
    para(doc, "This internship report has been submitted for examination with the approval of the following "
              "advisors and examiners.", justify=True)

    headers = ["Role", "Name", "Signature", "Date"]
    rows = [
        ["Intern", c1.STUDENT, "", ""],
        ["Internship Advisor (DBU)", "[Advisor Name]", "", ""],
        ["Internship Mentor (Company)", "[Mentor Name]", "", ""],
        ["Department Head / Coordinator", "[Coordinator Name]", "", ""],
    ]
    t = doc.add_table(rows=1 + len(rows), cols=4)
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = t.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run(r, size=11, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(cell, "E5E7EB")
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run(r, size=11)
    for i in range(len(t.rows)):
        t.cell(i, 0).width = Inches(2.3)
        t.cell(i, 1).width = Inches(2.0)
        t.cell(i, 2).width = Inches(1.2)
        t.cell(i, 3).width = Inches(1.2)


def build_prelim_page(doc, title_blocks, field_label=None):
    for tb in title_blocks:
        for kind, payload in tb:
            if kind == "h1":
                p = doc.add_paragraph(style="Heading 1")
                r = p.add_run(payload)
                set_run(r, size=14, bold=True)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(12)
            elif kind == "p":
                para(doc, payload)
    doc.add_page_break()


def render_toc_field(doc, code, placeholder=""):
    p = doc.add_paragraph()
    add_field(p, code, placeholder)
    return p


def field_note(doc):
    p = doc.add_paragraph()
    r = p.add_run("Note: This table of contents and the lists of tables and figures are Word fields. "
                  "In Microsoft Word, open the document and press Ctrl+A followed by F9 (or right-click "
                  "and choose \u201cUpdate Field\u201d) so that page numbers, figure numbers and table "
                  "numbers are generated automatically. The document is configured to update fields "
                  "automatically when opened.")
    set_run(r, size=9, italic=True, color=RGBColor(0x6B, 0x72, 0x80))


# ---------------------------------------------------------------- main
def main():
    doc = Document()
    setup_styles(doc)
    setattr(c1, "DURATION", "[Internship Start Date] \u2013 [Internship End Date], 2018 E.C.")
    setattr(c1, "SUBMIT", "[Month, Year] E.C.")

    # update fields on open
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.insert(0, uf)

    sec1 = setup_sections(doc)

    # -------- COVER (section 1) --------
    build_cover(doc)

    # -------- PRELIMINARY PAGES (section 2, roman) --------
    sec2 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_sections(doc)
    set_pg_fmt(sec2, "lowerRoman", start=1)
    sec2.footer.is_linked_to_previous = False
    fp = sec2.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp)

    build_declaration(doc)
    doc.add_page_break()

    # Acknowledgements
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("ACKNOWLEDGEMENTS")
    set_run(r, size=14, bold=True)
    p.paragraph_format.space_after = Pt(12)
    render_blocks(doc, c1.ACKNOWLEDGEMENTS)
    doc.add_page_break()

    # Executive Summary
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("EXECUTIVE SUMMARY")
    set_run(r, size=14, bold=True)
    p.paragraph_format.space_after = Pt(12)
    render_blocks(doc, c1.EXEC_SUMMARY)
    doc.add_page_break()

    # List of Tables
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("LIST OF TABLES")
    set_run(r, size=14, bold=True)
    p.paragraph_format.space_after = Pt(12)
    render_toc_field(doc, 'TOC \\h \\z \\c "Table"')
    doc.add_page_break()

    # List of Figures
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("LIST OF FIGURES")
    set_run(r, size=14, bold=True)
    p.paragraph_format.space_after = Pt(12)
    render_toc_field(doc, 'TOC \\h \\z \\c "Figure"')
    doc.add_page_break()

    # Table of Contents
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("TABLE OF CONTENTS")
    set_run(r, size=14, bold=True)
    p.paragraph_format.space_after = Pt(12)
    render_toc_field(doc, 'TOC \\o "1-3" \\h \\z \\u')
    field_note(doc)

    # -------- MAIN BODY (section 3, arabic) --------
    sec3 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_sections(doc)
    set_pg_fmt(sec3, "decimal", start=1)
    sec3.footer.is_linked_to_previous = False
    fp3 = sec3.footer.paragraphs[0]
    fp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp3)

    render_blocks(doc, c1.CHAPTER_1)
    render_blocks(doc, c1.CHAPTER_2)
    render_blocks(doc, c2.CHAPTER_3)
    render_blocks(doc, c2.CHAPTER_4)

    # References
    render_blocks(doc, c2.REFERENCES)

    # Appendices
    render_blocks(doc, ap.APPENDICES)

    # document properties
    doc.core_properties.title = ("Internship Report - Development of the FieldSync Offline-First "
                                 "Registration and Reporting System")
    doc.core_properties.author = c1.STUDENT
    doc.core_properties.subject = "Debre Berhan University - Internship Report (2018 E.C.)"

    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    main()